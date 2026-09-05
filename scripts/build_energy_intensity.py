#!/usr/bin/env python3
"""Writes src/data/learn_energy_intensity.json: world energy intensity
1965-2024 and the rate of every 25-year window inside it.

The tool measures energy intensity as primary energy over GDP at purchasing
power parity, and the World Bank's PPP series starts in 1990. To reach 1965
the World Bank level is carried backwards on the growth rates of the
Maddison Project Database's world total, which is also a PPP series. The
splice happens at 1990 and DATA.md records how far the two disagree over
the years they share.

Sources:
  Energy Institute Statistical Review 2026, total energy supply, world,
    1965-2024, exajoules. Read from the local workbook.
  World Bank NY.GDP.MKTP.PP.KD, world, 1990-2024, constant 2021
    international dollars.
  Maddison Project Database 2023, full dataset, country GDP per capita and
    population, restricted to the 151 countries with an unbroken record, so
    the growth rate compares the same economies in every year.

Run: python3 scripts/build_energy_intensity.py  [--refresh]
"""
import json
import sys
from pathlib import Path

import openpyxl

import build_data

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'src' / 'data' / 'learn_energy_intensity.json'
CACHE = ROOT / 'scripts' / '_maddison_cache.json'
MPD_URL = 'https://dataverse.nl/api/access/datafile/421302'
MPD_LOCAL = ROOT / 'scripts' / '_mpd2023.xlsx'

FIRST_YEAR, LAST_YEAR = 1965, 2024
SPLICE_YEAR = 1990          # where the World Bank series begins
WINDOW = 25                 # the span the brief asks the strip to show


def maddison_world() -> dict[int, float]:
    """World GDP by year from a sample of countries with an unbroken record.

    Coverage changes from year to year in the full dataset, and a changing
    sample would put steps in the growth rate. Restricting to countries
    present in every year of the span leaves 151 of them, 96.5% of world
    GDP in 1990, and lets each year compare with the last on like terms.
    """
    if CACHE.exists() and '--refresh' not in sys.argv:
        return {int(y): v for y, v in json.loads(CACHE.read_text()).items()}

    if not MPD_LOCAL.exists() or '--refresh' in sys.argv:
        print(f'downloading {MPD_URL}', file=sys.stderr)
        import urllib.request
        with urllib.request.urlopen(MPD_URL, timeout=600) as response:
            MPD_LOCAL.write_bytes(response.read())

    workbook = openpyxl.load_workbook(MPD_LOCAL, read_only=True)
    sheet = workbook['Full data']
    header = None
    by_year: dict[int, dict[str, float]] = {}
    for row in sheet.iter_rows(values_only=True):
        if header is None:
            header = [str(cell) for cell in row]
            continue
        record = dict(zip(header, row))
        year, code = record.get('year'), record.get('countrycode')
        gdppc, population = record.get('gdppc'), record.get('pop')
        if year is None or gdppc is None or population is None or code is None:
            continue
        if not FIRST_YEAR <= int(year) <= LAST_YEAR:
            continue
        by_year.setdefault(int(year), {})[str(code)] = float(gdppc) * float(population)

    years = sorted(by_year)
    unbroken = set.intersection(*[set(by_year[year]) for year in years])
    if len(unbroken) < 100:
        raise SystemExit(f'only {len(unbroken)} countries run unbroken; expected 100 or more')
    share = (sum(by_year[SPLICE_YEAR][c] for c in unbroken)
             / sum(by_year[SPLICE_YEAR].values()))
    if share < 0.9:
        raise SystemExit(f'the unbroken sample covers only {share:.1%} of {SPLICE_YEAR} GDP')
    print(f'Maddison: {len(unbroken)} countries unbroken {years[0]}-{years[-1]}, '
          f'{share:.1%} of {SPLICE_YEAR} world GDP', file=sys.stderr)

    world = {year: sum(by_year[year][c] for c in unbroken) for year in years}
    CACHE.write_text(json.dumps({str(y): v for y, v in world.items()}, indent=1) + '\n')
    return world


def spliced_gdp(worldbank: dict[int, float], maddison: dict[int, float]) -> dict[int, float]:
    """World Bank levels, carried back before 1990 on Maddison growth rates."""
    gdp = {year: value for year, value in worldbank.items() if year >= SPLICE_YEAR}
    anchor = gdp[SPLICE_YEAR]
    for year in range(SPLICE_YEAR - 1, FIRST_YEAR - 1, -1):
        gdp[year] = anchor * (maddison[year] / maddison[SPLICE_YEAR])
    return gdp


def windows(series: dict[int, float], span: int) -> list[dict]:
    out = []
    for start in range(FIRST_YEAR, LAST_YEAR - span + 1):
        end = start + span
        out.append({
            'from': start,
            'to': end,
            'value': round(build_data.cagr(series, start, end), 4),
        })
    return out


def percentile_of(rates: list[float], value: float) -> float:
    """Share of windows improving more slowly than `value`, in percent."""
    slower = sum(1 for rate in rates if rate > value)
    return round(100 * slower / len(rates), 1)


def main() -> None:
    workbook = openpyxl.load_workbook(build_data.WORKBOOK, read_only=True, data_only=True)
    energy_ej = build_data.read_sheet(workbook, 'Total Energy Supply (TES) -EJ')['Total World']
    energy = {year: value for year, value in energy_ej.items()
              if FIRST_YEAR <= year <= LAST_YEAR}

    cache = build_data.fetch_worldbank([])
    worldbank = {int(year): float(value) for year, value
                 in cache['world_gdp_ppp'].items() if value is not None}
    gdp = spliced_gdp(worldbank, maddison_world())

    # MJ per dollar: EJ is 1e18 J, GDP is in dollars, so EJ/GDP * 1e12 = MJ/$.
    intensity = {year: energy[year] / gdp[year] * 1e12 for year in
                 range(FIRST_YEAR, LAST_YEAR + 1)}

    maddison = maddison_world()
    overlap_maddison = build_data.cagr(maddison, SPLICE_YEAR, 2022)
    overlap_worldbank = build_data.cagr(worldbank, SPLICE_YEAR, 2022)

    # How much the splice itself decides. Maddison and the World Bank
    # disagree by 0.27 percentage points a year over the 32 years they share.
    # Carrying the level back on growth rates adjusted by that gap gives the
    # alternative record below, and the difference between the two whole-record
    # rates is the part of the answer the splice, rather than the data, chose.
    adjustment = (1 + overlap_worldbank / 100) / (1 + overlap_maddison / 100)
    alternative = dict(intensity)
    for year in range(SPLICE_YEAR - 1, FIRST_YEAR - 1, -1):
        gap = SPLICE_YEAR - year
        alternative[year] = energy[year] / (gdp[year] * adjustment ** -gap) * 1e12

    rate_windows = windows(intensity, WINDOW)
    rates = [w['value'] for w in rate_windows]
    years = list(range(FIRST_YEAR, LAST_YEAR + 1))

    payload = {
        'meta': {
            'generated_by': 'scripts/build_energy_intensity.py',
            'units': 'megajoules of primary energy per dollar of GDP, '
                     'constant 2021 international dollars',
            'sources': {
                'energy': 'Energy Institute Statistical Review 2026, '
                          'Total Energy Supply, Total World, EJ',
                'gdp_from_1990': 'World Bank NY.GDP.MKTP.PP.KD, WLD',
                'gdp_before_1990': 'Maddison Project Database 2023, full dataset, '
                                   '151 countries with an unbroken 1965-2022 record',
            },
            'splice': {
                'year': SPLICE_YEAR,
                'note': 'World Bank levels carried back on Maddison growth rates',
                'overlap_cagr_maddison': round(overlap_maddison, 4),
                'overlap_cagr_worldbank': round(overlap_worldbank, 4),
                'overlap_period': f'{SPLICE_YEAR}-2022',
            },
        },
        'series': [
            {
                'id': 'intensity',
                'label': 'Energy per dollar',
                'kind': 'history',
                'years': years,
                'values': [round(intensity[year], 4) for year in years],
            },
            {
                'id': 'energy',
                'label': 'Primary energy',
                'kind': 'history',
                'years': years,
                'values': [round(energy[year], 1) for year in years],
            },
            {
                'id': 'gdp',
                'label': 'World GDP',
                'kind': 'history',
                'years': years,
                'values': [round(gdp[year] / 1e12, 2) for year in years],
            },
        ],
        'windows': {
            'span': WINDOW,
            'rates': rate_windows,
            'fastest': min(rate_windows, key=lambda w: w['value']),
            'slowest': max(rate_windows, key=lambda w: w['value']),
        },
        'constants': {
            'firstYear': FIRST_YEAR,
            'lastYear': LAST_YEAR,
            'rates': {
                'longRecord': round(build_data.cagr(intensity, SPLICE_YEAR, LAST_YEAR), 4),
                'wholeRecord': round(build_data.cagr(intensity, FIRST_YEAR, LAST_YEAR), 4),
                'recentDecade': round(build_data.cagr(intensity, 2015, LAST_YEAR), 4),
            },
            'spliceSensitivity': {
                'wholeRecordRate': round(
                    build_data.cagr(alternative, FIRST_YEAR, LAST_YEAR), 4),
                'firstLevel': round(alternative[FIRST_YEAR], 4),
                'note': 'the same series with the pre-1990 growth rates adjusted to '
                        'match the World Bank over the years the two sources share',
            },
            'levels': {
                'first': round(intensity[FIRST_YEAR], 4),
                'splice': round(intensity[SPLICE_YEAR], 4),
                'last': round(intensity[LAST_YEAR], 4),
            },
            'energy': {
                'first': round(energy[FIRST_YEAR], 1),
                'last': round(energy[LAST_YEAR], 1),
                'growth': round(build_data.cagr(energy, FIRST_YEAR, LAST_YEAR), 4),
            },
            'gdp': {
                'growth': round(build_data.cagr(gdp, FIRST_YEAR, LAST_YEAR), 4),
            },
        },
    }
    OUT.write_text(json.dumps(payload, indent=1) + '\n')
    print(f'wrote {OUT.relative_to(ROOT)}: {intensity[FIRST_YEAR]:.2f} MJ/$ in {FIRST_YEAR}, '
          f'{intensity[LAST_YEAR]:.2f} in {LAST_YEAR}, {len(rate_windows)} windows of {WINDOW} '
          f'years, from {min(rates):.2f} to {max(rates):.2f} %/yr')


if __name__ == '__main__':
    main()
